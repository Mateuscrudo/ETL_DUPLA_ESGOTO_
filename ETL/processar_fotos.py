import re
import pandas as pd
from openpyxl import load_workbook

from .ler_bases import nk


def urls(v):
    if pd.isna(v):
        return []

    return [
        x.strip()
        for x in re.split(r"[\r\n]+", str(v))
        if x.strip().startswith("http")
    ]


def processar_fotos(roteiro_path, a):

    registros = []

    # =========================================================
    # ROTEIRO
    # =========================================================

    wb = load_workbook(
        roteiro_path,
        read_only=False,
        data_only=False
    )

    ws = wb["Planilha1"]

    headers = [
        ws.cell(1, c).value
        for c in range(1, ws.max_column + 1)
    ]

    coluna_foto = headers.index("Foto(s)") + 1
    coluna_matricula = headers.index("Matricula") + 1
    coluna_data = headers.index("Dt. Execução") + 1

    for linha in range(2, ws.max_row + 1):

        matricula = nk(
            ws.cell(linha, coluna_matricula).value
        )

        hyperlink = ws.cell(
            linha,
            coluna_foto
        ).hyperlink

        data = pd.to_datetime(
            str(ws.cell(linha, coluna_data).value),
            errors="coerce"
        )

        # Só adiciona se tiver matrícula E hyperlink
        if (
            matricula is not None
            and not pd.isna(matricula)
            and hyperlink
            and hyperlink.target
        ):

            registros.append(
                (
                    matricula,
                    "ROTEIRO",
                    "Galeria",
                    data,
                    hyperlink.target
                )
            )

    # =========================================================
    # APPDEEP
    # =========================================================

    colunas_foto = [
        coluna
        for coluna in a.columns
        if "foto" in coluna.lower()
        or "fotos" in coluna.lower()
    ]

    for _, row in a.iterrows():

        matricula = row["_mat"]

        # Proteção contra matrícula vazia
        if matricula is None or pd.isna(matricula):
            continue

        for coluna in colunas_foto:

            for url in urls(row[coluna]):

                registros.append(
                    (
                        matricula,
                        "APPDEEP",
                        coluna,
                        row["_dt"],
                        url
                    )
                )

    # =========================================================
    # DATAFRAME
    # =========================================================

    fotos = pd.DataFrame(
        registros,
        columns=[
            "Matrícula",
            "Origem",
            "Tipo",
            "Data do Registro",
            "URL"
        ]
    )

    if fotos.empty:
        print("⚠️ Nenhuma foto encontrada.")
        return fotos

    # =========================================================
    # LIMPEZA FINAL
    # =========================================================

    antes = len(fotos)

    fotos = fotos[
        fotos["Matrícula"].notna()
        & fotos["URL"].notna()
        & fotos["URL"].astype(str).str.strip().ne("")
    ].copy()

    descartadas = antes - len(fotos)

    print(
        f"📸 Fotos encontradas: {len(fotos):,}"
    )

    if descartadas:
        print(
            f"⚠️ Registros de fotos descartados: "
            f"{descartadas:,}"
        )

    # =========================================================
    # REMOVE DUPLICIDADES
    # =========================================================

    antes = len(fotos)

    fotos = (
        fotos
        .drop_duplicates(
            ["Matrícula", "URL"]
        )
        .sort_values(
            [
                "Matrícula",
                "Data do Registro",
                "Origem",
                "Tipo"
            ],
            na_position="last"
        )
        .reset_index(drop=True)
    )

    duplicadas = antes - len(fotos)

    if duplicadas:
        print(
            f"♻️ Fotos duplicadas removidas: "
            f"{duplicadas:,}"
        )

    return fotos